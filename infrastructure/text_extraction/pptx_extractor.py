from pptx import Presentation


def extract_pages_from_pptx(file_path: str) -> list[dict]:
    prs = Presentation(file_path)
    pages = []

    for i, slide in enumerate(prs.slides):
        text = "\n".join(
            shape.text for shape in slide.shapes if shape.has_text_frame
        )
        pages.append({"page_number": i + 1, "text": text})

    return pages